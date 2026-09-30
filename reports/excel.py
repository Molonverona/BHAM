"""
BHAM – Excel Report Generator (Enhanced v2.0)
Fogli:
- Network Topology: Albero gerarchico d'impianto (Host -> Interfaccia -> Bus -> Nodi)
- Modbus Devices: Censimento slave Modbus RTU e TCP
- BACnet Devices: Censimento controllori BACnet/IP e MS-TP
- KNX Devices: Censimento gateway e nodi KNXnet/IP
- IP Hosts: Host di rete rilevati tramite sniffing ARP L2
- Modbus Registers: Censimento registri scansionati (Smart Scan / Mappe BACS Help)
- BACnet Objects: Esplorazione analitica oggetti (AI, AO, AV, BI, BO, BV, MSI, MSO)
- Report Info: Metadati di collaudo e sommario statistico
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from data.state import state


# ── Palette Tipografica Industriale ──────────────────────────────────────────
_HDR_SLATE   = PatternFill("solid", fgColor="334155")   # Topology header
_HDR_CYAN    = PatternFill("solid", fgColor="0E7490")   # Modbus header
_HDR_PURPLE  = PatternFill("solid", fgColor="6D28D9")   # BACnet header
_HDR_ORANGE  = PatternFill("solid", fgColor="C2410C")   # KNX header
_HDR_GREEN   = PatternFill("solid", fgColor="065F46")   # Hosts header
_HDR_AMBER   = PatternFill("solid", fgColor="B45309")   # Modbus Registers header
_HDR_INDIGO  = PatternFill("solid", fgColor="4338CA")   # BACnet Objects header
_HDR_FONT    = Font(color="FFFFFF", bold=True)
_ALT_FILL    = PatternFill("solid", fgColor="F8FAFC")   # Alternate row zebra fill


def _write_sheet(
    wb: openpyxl.Workbook,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    fill: PatternFill,
) -> None:
    ws = wb.create_sheet(title)

    # Intestazioni colonna
    ws.append(headers)
    for col_idx, cell in enumerate(ws[1], start=1):
        cell.fill = fill
        cell.font = _HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Righe di dati
    for r_idx, row in enumerate(rows, start=2):
        ws.append(row)
        if r_idx % 2 == 0:
            for cell in ws[r_idx]:
                cell.fill = _ALT_FILL

    # Auto-width e allineamento
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = max((len(str(c.value or "")) for c in col_cells), default=8)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

    # Freeze riga di intestazione + filtri automatici
    ws.freeze_panes = ws["A2"]
    ws.auto_filter.ref = ws.dimensions


def generate_excel(filepath: str) -> str:
    """Genera report Excel arricchito (v2.0). Ritorna il filepath."""
    wb = openpyxl.Workbook()
    # Rimuove il foglio di default
    wb.remove(wb.active)

    # ── Sheet 1: Network Topology ────────────────────────────────────────────
    topo = state.get_topology()
    topo_headers = [
        "Node ID", "Livello", "Categoria", "Protocollo",
        "Nome / Etichetta", "Dettaglio Canale", "Parent Node ID", "Stato Operativo"
    ]
    topo_rows = []
    level_map = {"host": 1, "interface": 2, "bus": 3, "device": 4}
    for n in topo.get("nodes", []):
        cat = n.get("category", "device")
        lvl = level_map.get(cat, 4)
        topo_rows.append([
            n.get("id", ""),
            lvl,
            cat.upper(),
            n.get("protocol", "").upper(),
            n.get("label", ""),
            n.get("sublabel", ""),
            n.get("parent_id") or "—",
            n.get("status", "unknown").upper(),
        ])
    _write_sheet(wb, "Network Topology", topo_headers, topo_rows, _HDR_SLATE)

    # ── Sheet 2: Modbus Devices ──────────────────────────────────────────────
    mb_headers = [
        "Slave ID", "Protocol", "IP", "TCP Port",
        "Serial Port", "Baudrate", "Response ms", "Discovered At"
    ]
    mb_rows = []
    for d in state.modbus_devices.values():
        mb_rows.append([
            d.slave_id,
            d.protocol.value,
            d.ip or "",
            d.tcp_port if d.ip else "",
            d.serial_params.port if d.serial_params else "",
            d.serial_params.baudrate if d.serial_params else "",
            round(d.response_time_ms, 2) if d.response_time_ms else "",
            str(d.discovered_at)[:19],
        ])
    _write_sheet(wb, "Modbus Devices", mb_headers, mb_rows, _HDR_CYAN)

    # ── Sheet 3: BACnet Devices ──────────────────────────────────────────────
    bn_headers = [
        "Device ID", "Address", "Vendor ID", "Vendor Name",
        "Model Name", "FW Revision", "SW Version", "Discovered At"
    ]
    bn_rows = []
    for d in state.bacnet_devices.values():
        bn_rows.append([
            d.device_id,
            d.address,
            d.vendor_id or "",
            d.vendor_name or "",
            d.model_name or "",
            d.firmware_revision or "",
            d.application_software_version or "",
            str(d.discovered_at)[:19],
        ])
    _write_sheet(wb, "BACnet Devices", bn_headers, bn_rows, _HDR_PURPLE)

    # ── Sheet 4: KNX Devices ─────────────────────────────────────────────────
    knx_headers = [
        "Individual Address", "IP Address", "Port", "Device Name",
        "Serial Number", "MAC", "Medium", "Discovered At"
    ]
    knx_rows = []
    for k in state.knx_devices.values():
        knx_rows.append([
            k.individual_address,
            k.ip_address,
            k.port,
            k.device_name or "",
            k.serial_number or "",
            k.mac_address or "",
            k.medium,
            str(k.discovered_at)[:19],
        ])
    _write_sheet(wb, "KNX Devices", knx_headers, knx_rows, _HDR_ORANGE)

    # ── Sheet 5: IP Hosts (ARP) ──────────────────────────────────────────────
    ip_headers = [
        "IP Address", "MAC", "Hostname", "Open Ports",
        "Protocol Hints", "First Seen", "Last Seen"
    ]
    ip_rows = []
    for h in state.ip_hosts.values():
        ip_rows.append([
            h.ip,
            h.mac or "",
            h.hostname or "",
            ", ".join(str(p) for p in h.open_ports),
            ", ".join(p.value for p in h.protocol_hints),
            str(h.first_seen)[:19],
            str(h.last_seen)[:19],
        ])
    _write_sheet(wb, "IP Hosts", ip_headers, ip_rows, _HDR_GREEN)

    # ── Sheet 6: Modbus Smart Registers ──────────────────────────────────────
    reg_headers = [
        "Slave ID", "Protocol", "Endpoint Canale", "Indirizzo Registro",
        "Hex Raw", "Dec Int16", "Float32 IEEE 754", "Descrizione / Tag"
    ]
    reg_rows = []
    for d in sorted(state.modbus_devices.values(), key=lambda x: x.slave_id):
        endpoint = d.ip if d.ip else (d.serial_params.port if d.serial_params else "RS485")
        if d.ip and d.tcp_port:
            endpoint = f"{d.ip}:{d.tcp_port}"
        if d.registers:
            for reg_addr, rdata in sorted(d.registers.items(), key=lambda x: str(x[0])):
                if isinstance(rdata, dict):
                    hex_val = rdata.get("hex", "")
                    dec_val = rdata.get("int16", rdata.get("dec", ""))
                    float_val = rdata.get("float32", "")
                    desc = rdata.get("description", rdata.get("label", ""))
                else:
                    hex_val = hex(rdata) if isinstance(rdata, int) else ""
                    dec_val = str(rdata)
                    float_val = ""
                    desc = ""
                reg_rows.append([
                    d.slave_id,
                    d.protocol.value,
                    endpoint,
                    str(reg_addr),
                    hex_val,
                    dec_val,
                    float_val,
                    desc,
                ])
    _write_sheet(wb, "Modbus Registers", reg_headers, reg_rows, _HDR_AMBER)

    # ── Sheet 7: BACnet Objects Explorer ─────────────────────────────────────
    obj_headers = [
        "Device ID", "Costruttore", "Modello", "Object Identifier",
        "Tipo Oggetto", "Nome Oggetto", "Present Value", "Unità", "Status Flags"
    ]
    obj_rows = []
    for d in sorted(state.bacnet_devices.values(), key=lambda x: x.device_id):
        if d.object_list:
            for obj in d.object_list:
                obj_id = obj.get("object_identifier") or obj.get("id", "")
                obj_type = obj.get("object_type") or obj.get("type", "")
                obj_name = obj.get("object_name") or obj.get("name", "")
                pval = obj.get("present_value")
                pval_str = str(pval) if pval is not None else ""
                units = obj.get("units", "")
                flags = str(obj.get("status_flags", ""))
                obj_rows.append([
                    d.device_id,
                    d.vendor_name or "",
                    d.model_name or "",
                    obj_id,
                    obj_type,
                    obj_name,
                    pval_str,
                    units,
                    flags,
                ])
    _write_sheet(wb, "BACnet Objects", obj_headers, obj_rows, _HDR_INDIGO)

    # ── Sheet 8: Report Info / Metadata ──────────────────────────────────────
    ws_meta = wb.create_sheet("Report Info")
    ws_meta.append(["Parametro", "Valore"])
    ws_meta["A1"].font = Font(bold=True, color="FFFFFF")
    ws_meta["A1"].fill = _HDR_SLATE
    ws_meta["B1"].font = Font(bold=True, color="FFFFFF")
    ws_meta["B1"].fill = _HDR_SLATE

    ws_meta.append(["Generated by", "BHAM – BACS Help Auto Mapper v0.6.0"])
    ws_meta.append(["Generated at", str(datetime.now(timezone.utc))[:19] + " UTC"])
    cfg = state.session_config
    if cfg:
        ws_meta.append(["Sito / Impianto", cfg.site_name or "—"])
        ws_meta.append(["Porta Seriale RS485", f"{cfg.serial_port or '—'} @ {cfg.serial_baudrate}bps"])
        ws_meta.append(["NIC Rete Scansione", f"{cfg.scan_iface or '—'} ({cfg.scan_ip or '—'})"])
    ws_meta.append(["Nodi Topologia Totali", len(topo_rows)])
    ws_meta.append(["Dispositivi Modbus censiti", len(mb_rows)])
    ws_meta.append(["Controllori BACnet censiti", len(bn_rows)])
    ws_meta.append(["Nodi KNXnet/IP censiti", len(knx_rows)])
    ws_meta.append(["Host di Rete (ARP) censiti", len(ip_rows)])
    ws_meta.append(["Registri Modbus Mappati/Scansionati", len(reg_rows)])
    ws_meta.append(["Oggetti BACnet Censiti", len(obj_rows)])

    ws_meta.column_dimensions["A"].width = 35
    ws_meta.column_dimensions["B"].width = 50

    wb.save(filepath)
    return filepath
